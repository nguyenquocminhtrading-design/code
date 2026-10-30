import openpyxl, json
from pathlib import Path

lines = []

# 1. Per-period file 2020-01
wb = openpyxl.load_workbook('output/2020/2020-01/2020-01_cashew.xlsx', read_only=True, data_only=True)
lines.append('=== 2020-01_cashew.xlsx ===')
lines.append(f'Sheets: {wb.sheetnames}')
ws = wb['Tong_Quan_XNK']
rows = list(ws.iter_rows(values_only=True))
lines.append(f'Tong_Quan_XNK: {len(rows)-2} summary rows')
for r in rows[2:]:
    lines.append(f'  {r[2]} | {r[11]} | qty={r[5]} | val={r[6]} | file={str(r[12])[:45]}')
# Other rows per version sheet
for sn in wb.sheetnames:
    if sn.startswith(('XK_', 'NK_')) and sn not in ('XK_Theo_Nuoc', 'NK_Theo_Nuoc'):
        ws2 = wb[sn]
        rws = list(ws2.iter_rows(values_only=True))
        others = [r for r in rws[1:] if r and r[0] == 'Other']
        dbg = rws[1][8] if len(rws) > 1 and len(rws[1]) > 8 else ''
        lines.append(f'{sn}: {len(rws)-1} rows, Other={len(others)}, sample debug_file={dbg}')
wb.close()

# 2. All-years version files
lines.append('')
lines.append('=== All-years version files ===')
for suffix in ('ChinhThuc', 'DieuChinh', 'SoBo'):
    p = Path(f'output/all_years_cashew_{suffix}.xlsx')
    if p.exists():
        wbv = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        wsvq = wbv['Tong_Quan_XNK']
        n_sum = wsvq.max_row - 1
        wsvx = wbv['XK_Theo_Nuoc']
        wsvn = wbv['NK_Theo_Nuoc']
        lines.append(f'all_years_cashew_{suffix}.xlsx: sheets={wbv.sheetnames}, summary={n_sum}, XK={wsvx.max_row-1}, NK={wsvn.max_row-1}')
        wbv.close()
    else:
        lines.append(f'all_years_cashew_{suffix}.xlsx: MISSING')

# 3. JSON schema v3
with open('output/data/all_years_cashew.json', encoding='utf-8') as f:
    data = json.load(f)
lines.append('')
lines.append('=== JSON ===')
lines.append(f'schema_version: {data.get("schema_version")}')
recs = data.get('records', [])
with_vk = [r for r in recs if r.get('version_key')]
with_dbg = [r for r in recs if r.get('debug_source_file')]
lines.append(f'records: {len(recs)}, with version_key: {len(with_vk)}, with debug_file: {len(with_dbg)}')
p0101 = data['periods'].get('2020-01', {})
lines.append(f'2020-01 tong_xk_versions: {sorted((p0101.get("tong_xk_versions") or {}).keys())}')
vk_sample = next((r for r in recs if r.get('version_key')), None)
if vk_sample:
    lines.append(f'sample: period={vk_sample["period"]} country={vk_sample["country"]} vk={vk_sample["version_key"]} dbg={vk_sample.get("debug_source_file")}')

# 4. Reconciliation
rec_path = Path('output/qc/aggregate_reconciliation.json')
if rec_path.exists():
    rep = json.loads(rec_path.read_text(encoding='utf-8'))
    lines.append('')
    lines.append(f'=== QC reconciliation: status={rep.get("status")}, issues={rep.get("issue_count")} ===')
    for i in (rep.get('issues') or [])[:8]:
        lines.append(f'  - {i}')

Path('output/docs/verify_multiversion.txt').write_text('\n'.join(lines), encoding='utf-8')
print('Written output/docs/verify_multiversion.txt')
