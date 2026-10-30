import openpyxl
from pathlib import Path

xlsx_path = Path('output/2020/2020-01/2020-01_cashew.xlsx')
wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)

lines = []
lines.append('=' * 80)
lines.append('AUDIT: 2020-01 Excel Structure')
lines.append('=' * 80)
lines.append('')

# Check sheets
lines.append(f'Sheets: {wb.sheetnames}')
lines.append('')

# Check Tong_Quan_XNK
ws = wb['Tong_Quan_XNK']
rows = list(ws.iter_rows(values_only=True))
lines.append('=== Tong_Quan_XNK (Summary Rows) ===')
for i, row in enumerate(rows[:5]):
    lines.append(f'Row {i}: {row}')
lines.append(f'Total rows: {len(rows)}')
lines.append('')

# Check XK_Theo_Nuoc
ws = wb['XK_Theo_Nuoc']
rows = list(ws.iter_rows(values_only=True))
lines.append('=== XK_Theo_Nuoc (Country Rows) ===')
lines.append(f'Header: {rows[0]}')
lines.append(f'Total rows: {len(rows)}')
lines.append('')
lines.append('Sample rows (first 5):')
for i, row in enumerate(rows[1:6]):
    lines.append(f'Row {i+1}: {row}')

# Check Other rows
other_rows = [r for r in rows[1:] if r[0] == 'Other']
lines.append(f'')
lines.append(f'Other rows count: {len(other_rows)}')
for i, row in enumerate(other_rows):
    lines.append(f'  Other {i+1}: {row}')
lines.append('')

# Check Lich_Su_Phien_Ban
ws = wb['Lich_Su_Phien_Ban']
rows = list(ws.iter_rows(values_only=True))
lines.append('=== Lich_Su_Phien_Ban (Version History) ===')
lines.append(f'Header: {rows[0]}')
lines.append(f'Total rows: {len(rows)}')
for i, row in enumerate(rows[:10]):
    lines.append(f'Row {i}: {row}')

wb.close()

# Check PDF files
pdf_dir = Path('output/2020/2020-01/Xuat_Khau')
pdfs = sorted(pdf_dir.glob('*.pdf'))
lines.append('')
lines.append('=== PDF Files in Xuat_Khau ===')
for p in pdfs:
    lines.append(f'  {p.name}')

output_path = Path('output/docs/audit_2020_01_structure.txt')
output_path.parent.mkdir(parents=True, exist_ok=True)
output_path.write_text('\n'.join(lines), encoding='utf-8')
print(f'Report written to: {output_path}')
