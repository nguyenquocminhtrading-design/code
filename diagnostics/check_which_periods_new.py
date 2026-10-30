"""Kiểm tra kỳ nào đã chạy code mới (có sheet XK_ChinhThuc...) vs cấu trúc cũ."""
import openpyxl
from pathlib import Path

def check(xlsx):
    try:
        wb = openpyxl.load_workbook(str(xlsx), read_only=True, data_only=True)
        sheets = wb.sheetnames
        ws = wb['Tong_Quan_XNK']
        n_sum = max(0, ws.max_row - 2) if 'Tong_Quan_XNK' in sheets else 0
        wb.close()
        new_style = any(s.startswith(('XK_ChinhThuc', 'NK_ChinhThuc', 'XK_DieuChinh', 'NK_DieuChinh', 'XK_SoBo', 'NK_SoBo', 'XK_UocTinh', 'NK_UocTinh')) for s in sheets)
        return new_style, n_sum, len(sheets)
    except Exception as e:
        return None, 0, str(e)

master_new, master_old, flow_checked = [], [], 0
for year_dir in sorted(Path('output').iterdir()):
    if not (year_dir.is_dir() and year_dir.name.isdigit()):
        continue
    for pd_dir in sorted(year_dir.glob('????-??')):
        xlsx = pd_dir / f'{pd_dir.name}_cashew.xlsx'
        if not xlsx.exists():
            continue
        new_style, n_sum, extra = check(xlsx)
        (master_new if new_style else master_old).append(pd_dir.name)

print(f'MASTER FILES (new structure): {len(master_new)}')
for p in master_new:
    print(f'  NEW: {p}')
print(f'MASTER FILES (old structure): {len(master_old)}')
print('  Sample old:', master_old[:8], '...' if len(master_old) > 8 else '')

# Check flow folder workbook 2020-01
fw = Path('output/2020/2020-01/Nhap_Khau/2020-01_nhap_khau.xlsx')
if fw.exists():
    wb = openpyxl.load_workbook(str(fw), read_only=True)
    print(f'\nFlow workbook 2020-01 Nhap_Khau sheets: {wb.sheetnames}')
    wb.close()
