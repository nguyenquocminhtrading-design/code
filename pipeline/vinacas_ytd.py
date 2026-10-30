#!/usr/bin/env python3
"""Vinacas YTD builder — tự lũy kế YTD theo năm cho dữ liệu Vinacas (RCN import).

Mục đích
--------
Sinh dữ liệu CY (tháng) + YTD (lũy kế từ đầu năm) để dashboard visualize được
cả hai chế độ cho nguồn Vinacas:

* NĂM MONTHLY (2023+): mỗi kỳ là giá trị tháng → YTD = cumsum theo từng năm
  (mỗi năm reset, độc lập).
* NĂM ANNUAL (2018–2022): nguồn chỉ có duy nhất 1 kỳ YYYY-12 mang giá trị cả
  năm (đã cộng dồn sẵn) → quantity_ytd = quantity_month, is_annual = true,
  KHÔNG cộng kép.

Nguồn dữ liệu (--source)
------------------------
* json  : vinacas_records trong output/data/all_years_cashew.json (nhanh, đã chuẩn hóa).
* xlsx  : Vincass/Vinacas.xlsx (sheet gốc: Tên nước, Region, Mặt hàng, ĐVT,
          Lượng, Trị giá (USD), Fulldate, Export/Import).
* auto  : thử JSON trước, thiếu/thất bại thì fallback XLSX.

Đầu ra (--out-dir, mặc định output/)
------------------------------------
* vinacas_ytd_cashew.json : records per (period, country) + bản ghi tổng kỳ
  (__TOTAL__) + yearly_summary.
* vinacas_ytd_cashew.xlsx : 3 sheet CY_Monthly / YTD_Cumulative / Yearly_Summary.

CLI
---
    python pipeline/vinacas_ytd.py [--source auto|xlsx|json] [--out-dir output]
                                    [--update-all-years]

--update-all-years: ghi ngược quantity_ytd/value_ytd_usd vào vinacas_records của
output/data/all_years_cashew.json (TẠO BACKUP .bak TRƯỚC KHI GHI). Cờ này KHÔNG
được chạy mặc định — chỉ dùng khi được duyệt rõ ràng.

Quy tắc số liệu
---------------
* Không sửa số liệu nguồn, không làm tròn tùy tiện (giữ nguyên float).
* Nhiều dòng nguồn trùng (period, country) (vd 2 lô Guinea-Bissau cùng tháng)
  được HỢP NHẤT khi tổng hợp (cộng dồn) — từng giá trị gốc vẫn giữ nguyên.
* Ô trống (None) tham gia lũy kế với trọng số 0 nhưng giữ nguyên None ở cột CY.
"""
from __future__ import annotations

import argparse
import io
import json
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_ROOT = ROOT / "output"
ALL_YEARS_JSON = DEFAULT_OUTPUT_ROOT / "data" / "all_years_cashew.json"
VINACAS_XLSX = ROOT / "Vincass" / "Vinacas.xlsx"

TOTAL_KEY = "__TOTAL__"
DEFAULT_PRODUCT = "Raw Cashew Nut"
DEFAULT_UNIT = "Tấn"
SCHEMA_VERSION = 1


def force_utf8_console() -> None:
    """Windows console (cp1252) không in được tiếng Việt — ép UTF-8."""
    if hasattr(sys.stdout, "buffer"):
        sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "buffer"):
        sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")


# ──────────────────────────────── Nguồn dữ liệu ──────────────────────────────

def load_records_json(path: Path) -> list[dict]:
    """Đọc vinacas_records từ all_years_cashew.json → list row chuẩn."""
    with path.open("r", encoding="utf-8") as fh:
        payload = json.load(fh)
    rows = payload.get("vinacas_records") or []
    out: list[dict] = []
    for r in rows:
        period = str(r.get("period") or "")
        if len(period) != 7 or period[4] != "-":
            continue
        if not (period[:4].isdigit() and period[5:].isdigit()):
            continue
        country = r.get("country_report") or r.get("country") or r.get("country_raw") or ""
        if not str(country).strip():
            continue
        out.append({
            "period": period,
            "country": str(country).strip(),
            "product": r.get("product") or DEFAULT_PRODUCT,
            "unit": r.get("unit") or DEFAULT_UNIT,
            "quantity": r.get("quantity_month", r.get("quantity")),
            "value": r.get("value_month_usd", r.get("value_usd")),
            "flow": r.get("flow") or "import",
        })
    return out


_XLSX_COLUMN_MAP = {
    "tên nước data gốc": "country_raw",
    "ten nuoc data goc": "country_raw",
    "tên quốc gia trong report": "country_report",
    "ten quoc gia trong report": "country_report",
    "tên quốc gia trong báo cáo": "country_report",
    "region": "region",
    "khu vực": "region",
    "mặt hàng": "product",
    "mat hang": "product",
    "đvt": "unit",
    "dvt": "unit",
    "lượng": "quantity",
    "luong": "quantity",
    "khối lượng": "quantity",
    "trị giá (usd)": "value",
    "tri gia (usd)": "value",
    "giá trị (usd)": "value",
    "gia tri (usd)": "value",
    "value_usd": "value",
    "fulldate": "fulldate",
    "export/import": "direction",
    "exportimport": "direction",
}


def load_records_xlsx(path: Path) -> list[dict]:
    """Đọc Vincass/Vinacas.xlsx → list row chuẩn (ưu tiên pandas, fallback openpyxl)."""
    norm = lambda s: str(s or "").strip().lower()
    rows_raw: list[dict] = []
    try:
        import pandas as pd  # type: ignore

        frame = pd.read_excel(path, sheet_name=0)
        cols = {}
        for col in frame.columns:
            key = norm(col)
            cols[col] = _XLSX_COLUMN_MAP.get(key, key)
        frame = frame.rename(columns=cols)
        for _, row in frame.iterrows():
            rows_raw.append({k: row[k] for k in frame.columns})
    except ImportError:
        # Fallback thuần openpyxl (bundled runtime luôn có; phòng môi trường gọt)
        from openpyxl import load_workbook  # type: ignore

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if header is None:
            raise ValueError(f"Sheet trống: {path}")
        idx = {norm(h): i for i, h in enumerate(header) if h is not None}
        for values in it:
            rows_raw.append({
                mapped: values[i]
                for key, i in idx.items()
                if (mapped := _XLSX_COLUMN_MAP.get(key, key)) is not None
            })
        wb.close()

    out: list[dict] = []
    for r in rows_raw:
        country = r.get("country_report") or r.get("country_raw") or ""
        if not str(country).strip():
            continue
        direction = str(r.get("direction") or "Import").strip().lower()
        if direction and not direction.startswith("import"):
            continue  # nguồn Vinacas chỉ có Import RCN — bỏ dòng lạ nếu có
        fulldate = r.get("fulldate")
        period = None
        if fulldate is not None:
            try:
                import pandas as pd  # type: ignore

                ts = pd.to_datetime(fulldate, errors="coerce")
                if ts is not None and not (isinstance(ts, float) and pd.isna(ts)):
                    period = f"{ts.year:04d}-{ts.month:02d}"
            except Exception:
                period = None
        if period is None and isinstance(fulldate, str) and len(fulldate) >= 7:
            period = fulldate[:7].replace("/", "-")
        if period is None:
            continue
        out.append({
            "period": period,
            "country": str(country).strip(),
            "product": r.get("product") or DEFAULT_PRODUCT,
            "unit": r.get("unit") or DEFAULT_UNIT,
            "quantity": r.get("quantity"),
            "value": r.get("value"),
            "flow": "import",
        })
    return out


def load_rows(source: str) -> tuple[list[dict], str]:
    """auto: thử JSON (nhanh) trước, fallback XLSX."""
    if source in ("json", "auto"):
        if ALL_YEARS_JSON.exists():
            try:
                rows = load_records_json(ALL_YEARS_JSON)
                if rows:
                    return rows, "json"
            except (json.JSONDecodeError, OSError) as exc:
                if source == "json":
                    raise
                print(f"[warn] JSON lỗi ({exc}) → fallback XLSX", file=sys.stderr)
        elif source == "json":
            raise FileNotFoundError(f"Không tìm thấy {ALL_YEARS_JSON}")
    if VINACAS_XLSX.exists():
        rows = load_records_xlsx(VINACAS_XLSX)
        if rows:
            return rows, "xlsx"
    raise FileNotFoundError(
        f"Không đọc được dữ liệu Vinacas từ JSON ({ALL_YEARS_JSON}) lẫn XLSX ({VINACAS_XLSX})"
    )


# ─────────────────────────────── Tính toán YTD ───────────────────────────────

def _num(v) -> float | None:
    """Số an toàn: None/NaN → None; chuỗi số → float."""
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    if f != f:  # NaN
        return None
    return f


def detect_annual_years(rows: list[dict]) -> dict[int, bool]:
    """Năm ANNUAL = toàn bộ dòng của năm chỉ nằm ở đúng 1 kỳ (vd 2019 chỉ có 2019-12)."""
    year_periods: dict[int, set[str]] = defaultdict(set)
    for r in rows:
        year_periods[int(r["period"][:4])].add(r["period"])
    return {year: len(periods) == 1 for year, periods in year_periods.items()}


def aggregate_per_period_country(rows: list[dict]) -> dict[tuple[str, str], dict]:
    """Hợp nhất các dòng nguồn trùng (period, country) — cộng dồn, giữ nguyên gốc."""
    agg: dict[tuple[str, str], dict] = {}
    for r in rows:
        key = (r["period"], r["country"])
        slot = agg.get(key)
        if slot is None:
            agg[key] = {
                "period": r["period"],
                "country": r["country"],
                "product": r["product"],
                "unit": r["unit"],
                "flow": r["flow"],
                "quantity": _num(r["quantity"]),
                "value": _num(r["value"]),
            }
            continue
        q, v = _num(r["quantity"]), _num(r["value"])
        if q is not None:
            slot["quantity"] = q if slot["quantity"] is None else slot["quantity"] + q
        if v is not None:
            slot["value"] = v if slot["value"] is None else slot["value"] + v
    return agg


def compute_ytd(agg: dict[tuple[str, str], dict], annual_years: dict[int, bool]) -> dict[tuple[str, str], dict]:
    """YTD tự lũy kế: mỗi NĂM độc lập, cumsum theo period. ANNUAL → YTD = CY."""
    by_series: dict[tuple[str, int], list[dict]] = defaultdict(list)
    for (period, country), slot in agg.items():
        by_series[(country, int(period[:4]))].append(slot)

    result: dict[tuple[str, str], dict] = {}
    for (country, year), slots in by_series.items():
        slots.sort(key=lambda s: s["period"])
        if annual_years.get(year, False):
            # Năm báo cáo gốc đã cộng dồn — YTD giữ nguyên giá trị kỳ (không cộng kép)
            for s in slots:
                result[(s["period"], country)] = {
                    **s, "quantity_ytd": s["quantity"], "value_ytd": s["value"],
                }
            continue
        run_q = run_v = 0.0
        for s in slots:
            if s["quantity"] is not None:
                run_q += s["quantity"]
            if s["value"] is not None:
                run_v += s["value"]
            result[(s["period"], country)] = {
                **s, "quantity_ytd": run_q, "value_ytd": run_v,
            }
    return result


def build_totals(agg: dict[tuple[str, str], dict], annual_years: dict[int, bool]) -> dict[tuple[str, str], dict]:
    """Tổng mỗi kỳ (__TOTAL__) = cộng mọi quốc gia, rồi tự lũy kế YTD theo năm."""
    per_period: dict[str, dict] = {}
    for (period, _country), slot in agg.items():
        slot_t = per_period.get(period)
        if slot_t is None:
            per_period[period] = {
                "period": period, "country": TOTAL_KEY,
                "product": slot["product"], "unit": slot["unit"], "flow": slot["flow"],
                "quantity": slot["quantity"], "value": slot["value"],
            }
            continue
        if slot["quantity"] is not None:
            slot_t["quantity"] = slot["quantity"] if slot_t["quantity"] is None else slot_t["quantity"] + slot["quantity"]
        if slot["value"] is not None:
            slot_t["value"] = slot["value"] if slot_t["value"] is None else slot_t["value"] + slot["value"]
    total_agg = {(s["period"], TOTAL_KEY): s for s in per_period.values()}
    return compute_ytd(total_agg, annual_years)


def apply_yoy(rows: list[dict]) -> None:
    """YoY của YTD: so cùng kỳ (cùng month) năm trước. Thiếu/mẫu số 0 → None."""
    index = {(int(r["period"][:4]), int(r["period"][5:7]), r["country"]): r for r in rows}
    for r in rows:
        year, month, country = int(r["period"][:4]), int(r["period"][5:7]), r["country"]
        prev = index.get((year - 1, month, country))
        for cur_key, prev_key, out_key in (
            ("quantity_ytd", "quantity_ytd", "ytd_qty_yoy_pct"),
            ("value_ytd", "value_ytd", "ytd_val_yoy_pct"),
        ):
            cur, pv = r.get(cur_key), prev.get(prev_key) if prev else None
            r[out_key] = (cur - pv) / pv * 100.0 if (cur is not None and pv not in (None, 0)) else None


def build_payload(rows: list[dict], source: str) -> dict:
    annual_years = detect_annual_years(rows)
    agg = aggregate_per_period_country(rows)
    country_ytd = compute_ytd(agg, annual_years)
    total_ytd = build_totals(agg, annual_years)

    records: list[dict] = []
    for key, slot in {**country_ytd, **total_ytd}.items():
        records.append({
            "period": slot["period"],
            "country": slot["country"],
            "product": slot["product"],
            "unit": slot["unit"],
            "flow": slot["flow"] or "import",
            "quantity_month": slot["quantity"],
            "value_month_usd": slot["value"],
            "quantity_ytd": slot["quantity_ytd"],
            "value_ytd_usd": slot["value_ytd"],
            "is_annual": bool(annual_years.get(int(slot["period"][:4]), False)),
            "ytd_qty_yoy_pct": None,
            "ytd_val_yoy_pct": None,
        })
    records.sort(key=lambda r: (r["period"], r["country"]))
    apply_yoy(records)

    yearly_summary: list[dict] = []
    for year in sorted({int(r["period"][:4]) for r in records}):
        yr_rows = [r for r in records if int(r["period"][:4]) == year and r["country"] == TOTAL_KEY]
        last = max(yr_rows, key=lambda r: r["period"])
        yearly_summary.append({
            "year": year,
            "cy_qty_total": last["quantity_ytd"] if annual_years[year] else sum(
                (r["quantity_month"] or 0) for r in yr_rows),
            "cy_val_total": last["value_ytd_usd"] if annual_years[year] else sum(
                (r["value_month_usd"] or 0) for r in yr_rows),
            "ytd_qty_last": last["quantity_ytd"],
            "ytd_val_last": last["value_ytd_usd"],
            "is_annual": bool(annual_years[year]),
        })

    n_dup_pairs = len(rows) - len(agg)
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": source,
        "meta": {
            "rows_source": len(rows),
            "rows_emitted": len(records),
            "consolidated_duplicate_pairs": max(n_dup_pairs, 0),
            "annual_years": sorted(y for y, a in annual_years.items() if a),
            "monthly_years": sorted(y for y, a in annual_years.items() if not a),
            "ytd_rule": "monthly=cumsum theo năm (reset đầu năm); annual(2018-2022, 1 kỳ/năm)=giữ nguyên giá trị kỳ",
            "null_value_rows": sum(1 for r in rows if _num(r["value"]) is None),
        },
        "records": records,
        "yearly_summary": yearly_summary,
    }


# ────────────────────────────────── Xuất file ────────────────────────────────

def write_json(payload: dict, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "vinacas_ytd_cashew.json"
    with path.open("w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    return path


def _records_frame(payload: dict, kind: str) -> list[dict]:
    rows = []
    for r in payload["records"]:
        year, month = int(r["period"][:4]), int(r["period"][5:7])
        base = {"Kỳ": r["period"], "Năm": year, "Tháng": month, "Quốc gia": r["country"],
                "Mặt hàng": r["product"], "ĐVT": r["unit"], "Luồng": r["flow"]}
        if kind == "cy":
            rows.append({**base, "CY Lượng (Tấn)": r["quantity_month"],
                         "CY Trị giá (USD)": r["value_month_usd"]})
        else:
            rows.append({**base, "YTD Lượng (Tấn)": r["quantity_ytd"],
                         "YTD Trị giá (USD)": r["value_ytd_usd"],
                         "Năm cộng dồn": r["is_annual"],
                         "YTD Lượng YoY%": r["ytd_qty_yoy_pct"],
                         "YTD Trị giá YoY%": r["ytd_val_yoy_pct"]})
    return rows


def write_excel(payload: dict, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "vinacas_ytd_cashew.xlsx"
    summary_rows = [{
        "Năm": s["year"], "CY Tổng Lượng (Tấn)": s["cy_qty_total"],
        "CY Tổng Trị giá (USD)": s["cy_val_total"],
        "YTD cuối năm Lượng (Tấn)": s["ytd_qty_last"],
        "YTD cuối năm Trị giá (USD)": s["ytd_val_last"],
        "Chế độ": "Annual (báo cáo năm)" if s["is_annual"] else "Monthly (theo tháng)",
    } for s in payload["yearly_summary"]]
    try:
        import pandas as pd  # type: ignore

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            pd.DataFrame(_records_frame(payload, "cy")).to_excel(
                writer, sheet_name="CY_Monthly", index=False)
            pd.DataFrame(_records_frame(payload, "ytd")).to_excel(
                writer, sheet_name="YTD_Cumulative", index=False)
            pd.DataFrame(summary_rows).to_excel(
                writer, sheet_name="Yearly_Summary", index=False)
    except ImportError:
        from openpyxl import Workbook  # type: ignore

        wb = Workbook()
        for sheet, rows in (("CY_Monthly", _records_frame(payload, "cy")),
                            ("YTD_Cumulative", _records_frame(payload, "ytd")),
                            ("Yearly_Summary", summary_rows)):
            ws = wb.create_sheet(sheet)
            headers = list(rows[0].keys()) if rows else []
            ws.append(headers)
            for row in rows:
                ws.append([row[h] for h in headers])
        del wb["Sheet"]
        wb.save(path)
    return path


def update_all_years(payload: dict, path: Path) -> None:
    """Ghi ngược YTD vào vinacas_records của all_years_cashew.json (cờ tùy chọn)."""
    ytd_by_key = {(r["period"], r["country"]): (r["quantity_ytd"], r["value_ytd_usd"])
                  for r in payload["records"] if r["country"] != TOTAL_KEY}
    backup = path.with_suffix(".json.bak")
    backup.write_bytes(path.read_bytes())
    with path.open("r", encoding="utf-8") as fh:
        doc = json.load(fh)
    hit = 0
    for r in doc.get("vinacas_records") or []:
        country = r.get("country_report") or r.get("country") or r.get("country_raw") or ""
        ytd = ytd_by_key.get((str(r.get("period") or ""), str(country).strip()))
        if not ytd:
            continue
        r["quantity_ytd"], r["value_ytd_usd"] = ytd
        r["source_note"] = "Vinacas · YTD tự lũy kế theo năm (pipeline/vinacas_ytd.py)"
        hit += 1
    with path.open("w", encoding="utf-8") as fh:
        json.dump(doc, fh, ensure_ascii=False, indent=2)
    print(f"[update-all-years] backup → {backup.name}; đã cập nhật {hit}/{len(doc.get('vinacas_records') or [])} bản ghi")


# ─────────────────────────────────── Console ─────────────────────────────────

def print_yearly_table(payload: dict) -> None:
    print("=" * 96)
    print("VINACAS RCN IMPORT — TỔNG HỢP THEO NĂM (CY = tháng; YTD = tự lũy kế từ đầu năm)")
    print("=" * 96)
    header = f"{'Năm':<6}{'CY Lượng (Tấn)':>18}{'CY Trị giá (USD)':>20}{'YTD cuối Lượng':>18}{'YTD cuối Trị giá':>20}  {'Chế độ':<20}"
    print(header)
    print("-" * 96)
    for s in payload["yearly_summary"]:
        mode = "Annual (báo cáo năm)" if s["is_annual"] else "Monthly (theo tháng)"
        print(f"{s['year']:<6}{s['cy_qty_total']:>18,.2f}{s['cy_val_total']:>20,.2f}"
              f"{s['ytd_qty_last']:>18,.2f}{s['ytd_val_last']:>20,.2f}  {mode:<20}")
    print("=" * 96)


def print_selfcheck(payload: dict) -> None:
    """Đối chiếu bắt buộc trước khi báo cáo: YTD cuối năm 2024 & 2025."""
    print("SELF-CHECK:")
    for s in payload["yearly_summary"]:
        if s["year"] in (2024, 2025):
            print(f"  YTD cuối năm {s['year']}: Lượng = {s['ytd_qty_last']:,.2f} Tấn | "
                  f"Trị giá = {s['ytd_val_last']:,.2f} USD")


# ──────────────────────────────────── CLI ────────────────────────────────────

def main(argv: list[str] | None = None) -> int:
    force_utf8_console()
    parser = argparse.ArgumentParser(
        description="Tính YTD tự lũy kế theo năm cho dữ liệu Vinacas (RCN import).")
    parser.add_argument("--source", choices=("auto", "xlsx", "json"), default="auto",
                        help="Nguồn dữ liệu: auto (JSON trước, fallback XLSX) | json | xlsx")
    parser.add_argument("--out-dir", default=str(DEFAULT_OUTPUT_ROOT),
                        help="Thư mục xuất kết quả (mặc định: output/)")
    parser.add_argument("--update-all-years", action="store_true",
                        help="GHI NGƯỢC YTD vào output/data/all_years_cashew.json (tạo .bak; chỉ chạy khi được duyệt)")
    args = parser.parse_args(argv)

    rows, source = load_rows(args.source)
    print(f"[source] {source}: {len(rows)} dòng nguồn "
          f"({'JSON' if source == 'json' else 'Excel'} → {ALL_YEARS_JSON if source == 'json' else VINACAS_XLSX})")

    payload = build_payload(rows, source)
    meta = payload["meta"]
    print(f"[agg] {meta['rows_source']} dòng nguồn → {meta['rows_emitted']} bản ghi "
          f"(hợp nhất {meta['consolidated_duplicate_pairs']} cặp trùng period+country; "
          f"dòng trị giá trống: {meta['null_value_rows']})")
    print(f"[rule] năm ANNUAL (1 kỳ/năm, không cộng kép): {meta['annual_years']}")
    print(f"[rule] năm MONTHLY (cumsum theo tháng): {meta['monthly_years']}")

    json_path = write_json(payload, Path(args.out_dir))
    xlsx_path = write_excel(payload, Path(args.out_dir))
    print_yearly_table(payload)
    print_selfcheck(payload)
    print(f"[out] {json_path}")
    print(f"[out] {xlsx_path}")

    if args.update_all_years:
        update_all_years(payload, ALL_YEARS_JSON)
    return 0


if __name__ == "__main__":
    sys.exit(main())
